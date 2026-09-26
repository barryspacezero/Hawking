"""Extract readable text blocks from uploaded documents."""

from __future__ import annotations

from dataclasses import dataclass
import io
import json
import os
import shutil
import subprocess
import tempfile
from typing import List, Optional

import fitz  # PyMuPDF

SUPPORTED_EXTENSIONS = frozenset({
    "pdf",
    "txt",
    "text",
    "md",
    "markdown",
    "csv",
    "log",
    "docx",
    "doc",
    "epub",
    "rtf",
    "odt",
    "html",
    "htm",
})

EXTENSION_ALIASES = {
    "text": "txt",
    "markdown": "md",
    "htm": "html",
}


@dataclass
class ParsedBlock:
    block_index: int
    page_number: Optional[int]
    text: str
    text_spans: Optional[str] = None


@dataclass
class ParseResult:
    blocks: List[ParsedBlock]
    canonical_type: str
    page_count: Optional[int] = None


def normalize_extension(filename: str) -> str:
    if "." not in filename:
        return ""
    ext = filename.rsplit(".", 1)[-1].lower()
    return EXTENSION_ALIASES.get(ext, ext)


def is_supported_extension(ext: str) -> bool:
    return ext in SUPPORTED_EXTENSIONS


def supported_formats_message() -> str:
    formats = sorted({EXTENSION_ALIASES.get(ext, ext) for ext in SUPPORTED_EXTENSIONS})
    return ", ".join(f".{fmt}" for fmt in formats)


def _paragraph_blocks(text: str) -> List[ParsedBlock]:
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs and text.strip():
        paragraphs = [line.strip() for line in text.split("\n") if line.strip()]
    if not paragraphs:
        raise ValueError("No extractable text found in file.")
    return [
        ParsedBlock(block_index=idx, page_number=None, text=paragraph)
        for idx, paragraph in enumerate(paragraphs)
    ]


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode text file. Try saving as UTF-8.")


def _extract_pdf_page_spans(page) -> str:
    rect = page.rect
    words = page.get_text("words")
    word_spans = [
        {"word": word[4], "x0": word[0], "y0": word[1], "x1": word[2], "y1": word[3]}
        for word in words
        if word[4].strip()
    ]
    return json.dumps({
        "page_width": rect.width,
        "page_height": rect.height,
        "words": word_spans,
    })


def _extract_pdf(content: bytes) -> ParseResult:
    pdf_doc = fitz.open(stream=content, filetype="pdf")
    page_count = len(pdf_doc)
    blocks: List[ParsedBlock] = []
    block_idx = 0
    for page_num, page in enumerate(pdf_doc):
        text = page.get_text("text").strip()
        if text:
            blocks.append(
                ParsedBlock(
                    block_index=block_idx,
                    page_number=page_num + 1,
                    text=text,
                    text_spans=_extract_pdf_page_spans(page),
                )
            )
            block_idx += 1
    pdf_doc.close()
    if not blocks:
        raise ValueError("No extractable text found. This may be an image-only or encrypted PDF.")
    return ParseResult(blocks=blocks, canonical_type="pdf", page_count=page_count)


def _extract_epub(content: bytes) -> ParseResult:
    epub_doc = fitz.open(stream=content, filetype="epub")
    page_count = len(epub_doc)
    blocks: List[ParsedBlock] = []
    block_idx = 0
    for page_num, page in enumerate(epub_doc):
        text = page.get_text("text").strip()
        if text:
            blocks.append(
                ParsedBlock(
                    block_index=block_idx,
                    page_number=page_num + 1,
                    text=text,
                )
            )
            block_idx += 1
    epub_doc.close()
    if not blocks:
        raise ValueError("No extractable text found in EPUB file.")
    return ParseResult(blocks=blocks, canonical_type="epub", page_count=page_count)


def _extract_docx(content: bytes) -> ParseResult:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts: List[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    if not parts:
        raise ValueError("No extractable text found in DOCX file.")
    return ParseResult(blocks=_paragraph_blocks("\n\n".join(parts)), canonical_type="docx")


def _extract_rtf(content: bytes) -> ParseResult:
    from striprtf.striprtf import rtf_to_text

    text = rtf_to_text(content.decode("latin-1", errors="replace"))
    return ParseResult(blocks=_paragraph_blocks(text), canonical_type="rtf")


def _extract_odt(content: bytes) -> ParseResult:
    from odf.opendocument import load
    from odf import teletype

    doc = load(io.BytesIO(content))
    text = teletype.extractText(doc)
    return ParseResult(blocks=_paragraph_blocks(text), canonical_type="odt")


def _extract_html(content: bytes) -> ParseResult:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(content, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text(separator="\n\n").strip()
    return ParseResult(blocks=_paragraph_blocks(text), canonical_type="html")


def _extract_doc_with_antiword(content: bytes) -> Optional[str]:
    if not shutil.which("antiword"):
        return None
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name
    try:
        result = subprocess.run(
            ["antiword", tmp_path],
            capture_output=True,
            text=True,
            timeout=30,
            check=False,
        )
    finally:
        os.unlink(tmp_path)
    if result.returncode == 0 and result.stdout.strip():
        return result.stdout
    return None


def _extract_doc_with_libreoffice(content: bytes) -> Optional[str]:
    if not shutil.which("soffice"):
        return None
    with tempfile.TemporaryDirectory() as tmp_dir:
        input_path = os.path.join(tmp_dir, "upload.doc")
        with open(input_path, "wb") as handle:
            handle.write(content)
        result = subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "txt:Text",
                "--outdir",
                tmp_dir,
                input_path,
            ],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        if result.returncode != 0:
            return None
        output_path = os.path.join(tmp_dir, "upload.txt")
        if not os.path.exists(output_path):
            return None
        with open(output_path, "r", encoding="utf-8", errors="replace") as handle:
            text = handle.read().strip()
        return text or None


def _extract_doc(content: bytes) -> ParseResult:
    text = _extract_doc_with_antiword(content) or _extract_doc_with_libreoffice(content)
    if not text:
        raise ValueError(
            "Failed to parse .doc file. Convert to .docx or .pdf, or install antiword/LibreOffice on the server."
        )
    return ParseResult(blocks=_paragraph_blocks(text), canonical_type="doc")


def parse_document(filename: str, content: bytes) -> ParseResult:
    ext = normalize_extension(filename)
    if not is_supported_extension(ext):
        raise ValueError(f"Unsupported file type. Supported formats: {supported_formats_message()}.")

    if ext == "pdf":
        return _extract_pdf(content)
    if ext == "epub":
        return _extract_epub(content)
    if ext == "docx":
        return _extract_docx(content)
    if ext == "doc":
        return _extract_doc(content)
    if ext == "rtf":
        return _extract_rtf(content)
    if ext == "odt":
        return _extract_odt(content)
    if ext in {"html", "htm"}:
        return _extract_html(content)
    if ext in {"txt", "text", "md", "markdown", "csv", "log"}:
        return ParseResult(
            blocks=_paragraph_blocks(_decode_text(content)),
            canonical_type=EXTENSION_ALIASES.get(ext, ext),
        )

    raise ValueError(f"Unsupported file type. Supported formats: {supported_formats_message()}.")
